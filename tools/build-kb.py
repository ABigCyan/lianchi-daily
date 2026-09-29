"""
把《健身Excel超级套表》里的文字整理成助手可以检索的知识库 www/js/data-kb.js。
用法：python3 tools/build-kb.py <套表.xlsx>
每一行单元格合成一条，带“表号 + 行号 + 列号”，助手回答时按这个标出处。
表1-15 的饮食表结构相同，内容相同的行只保留一条，并记下出现在哪几张表。
"""
import json, re, sys
import openpyxl

src = sys.argv[1]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
# 这些表里的数字是查表数据（营养率、有氧消耗、价格、力量公式），要保留；其余表只保留文字
KEEP_NUM = {'16', '19', '20', '25'}
SKIP = {'26', '27', '29', '30'}  # 只有图片
rows, seen = [], {}
sheets = []
for name in wb.sheetnames:
    m = re.match(r'^(\d+)(.*)$', name)
    no = m.group(1) if m else '0'
    title = (m.group(2) if m else name).strip()
    if no in SKIP:
        continue
    sheets.append({'s': no, 'n': title})
    ws = wb[name]
    for r, row in enumerate(ws.iter_rows(), start=1):
        cells = []
        for c in row:
            v = c.value
            if v is None or not hasattr(c, 'column_letter'):
                continue
            if isinstance(v, str):
                v = re.sub(r'\s+', ' ', v.replace('\n', ' / ')).strip()
                if not v or v in {'*', '-'}:
                    continue
            elif isinstance(v, (int, float)) and no in KEEP_NUM:
                v = ('%.3f' % v).rstrip('0').rstrip('.') if isinstance(v, float) else str(v)
            else:
                continue
            cells.append('%s:%s' % (c.column_letter, v))
        if not cells:
            continue
        text = ' ｜ '.join(cells)
        if not re.search(r'[一-鿿]', text):
            continue
        # 饮食表 1-15：同样文字只存一条
        if no.isdigit() and 1 <= int(no) <= 15:
            key = re.sub(r'^[A-Z]+:', '', text)
            key = re.sub(r'(?<=｜ )[A-Z]+:', '', key)
            if key in seen:
                seen[key]['o'].append(no)
                continue
            item = {'s': no, 'r': r, 't': text, 'o': []}
            seen[key] = item
            rows.append(item)
        else:
            rows.append({'s': no, 'r': r, 't': text})
out = []
for it in rows:
    d = {'s': it['s'], 'r': it['r'], 't': it['t']}
    if it.get('o'):
        d['o'] = ','.join(it['o'])
    out.append(d)
js = ('/* 自动生成：python3 tools/build-kb.py 套表.xlsx —— 《健身Excel超级套表》（B站好人松松，原文件注明可任意分享）的文字内容，供助手检索并标注出处 */\n'
      'window.KB = ' + json.dumps({'sheets': sheets, 'rows': out}, ensure_ascii=False, separators=(',', ':')) + ';\n')
open('www/js/data-kb.js', 'w').write(js)
print(len(out), 'rows,', len(js.encode()) // 1024, 'KB')
