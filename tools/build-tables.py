"""
从《健身Excel超级套表》原表生成 App 用的查表数据，保证和原表逐行一致：
  - 表16 有氧热量消耗：全部 64 个项目（方法二 E80-E143）、跑步速度表、体重修正系数（从 F-T 列公式里读出）
  - 表19 日常食物营养率：全部碳水、蛋白质、含蛋白质的混合物
  - 表26/27 拉伸图谱：14 张拉伸图（压缩成 1000px JPEG）
  - test/fixtures/table16.json：表16 每个格子的原表计算结果，测试时逐格对照
用法：python3 tools/build-tables.py <套表.xlsx>
"""
import io, json, re, sys, zipfile, posixpath
import openpyxl
from PIL import Image

SRC = sys.argv[1]
wv = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
wf = openpyxl.load_workbook(SRC, read_only=True)
WEIGHTS = [50, 55, 60, 65, 70, 75, 80, 85, 90, 95, 100, 105, 110, 115, 120]
COLS = 'FGHIJKLMNOPQRST'


def rows_of(ws, lo, hi):
    out = {}
    for r in ws.iter_rows(min_row=lo, max_row=hi):
        for c in r:
            if hasattr(c, 'column_letter') and c.value is not None:
                out.setdefault(c.row, {})[c.column_letter] = c.value
    return out


# ---------- 表16 ----------
V16 = rows_of(wv['16有氧热量消耗'], 14, 143)
F16 = rows_of(wf['16有氧热量消耗'], 14, 143)
factor_of_row = {}
for row, d in F16.items():
    fs = []
    for col, w in zip(COLS, WEIGHTS):
        m = re.match(r'=ROUND\(\$E\d+\*(\d+)(?:\*([\d.]+))?,-1\)', str(d.get(col, '')))
        if m:
            fs.append(float(m.group(2) or 1))
    if len(fs) == 15:
        factor_of_row[row] = tuple(fs)
kinds = sorted(set(factor_of_row.values()))
flat = [k for k in kinds if k[6] == 1.0][0]
normal = [k for k in kinds if k[6] != 1.0][0]

def clean(s):
    return re.sub(r'\s+', '', str(s).split('\n')[0]).replace('*', '')

LEGACY = {  # 旧版本保存过的名字保持不变，免得用户资料里的有氧项目对不上
    ('平地走', '每走一小时'): '快走/散步（每小时）', ('跳操跟练', '中等强度'): '跳操（中等强度）',
    ('跳绳', '100-120次/分钟'): '跳绳 100-120 次/分',
}
SWIM = {0: '游泳（慢）', 1: '游泳（中）', 2: '游泳（快）'}
cardio, run_table, expect = [], [], {}
cat, swim_i = '', 0
for row in range(80, 144):
    d = V16.get(row, {})
    if 'C' in d:
        cat = clean(d['C'])
    item = str(d.get('D', '')).strip()
    coef = d.get('E')
    if not isinstance(coef, (int, float)):
        continue
    expect[row] = [d.get(c) for c in COLS]
    isflat = factor_of_row.get(row) == flat
    if cat == '跑步':
        run_table.append([float(item), coef])
        continue
    if (cat, item) in LEGACY:
        label = LEGACY[(cat, item)]
    elif cat.startswith('爬坡走'):
        label = '跑步机坡度 ' + item.replace('跑步机坡度', '')
    elif cat.startswith('户外骑行'):
        label = f'骑行 {item} km/h'
    elif cat.startswith('游泳'):
        label = SWIM[swim_i] + f' {float(item):.1f} 米/秒'; swim_i += 1
    elif cat == '室内单车':
        label = '动感单车 ' + item.replace('功率', '')
    elif cat == '划船机':
        label = '划船机 ' + item.replace('功率', '')
    elif cat.startswith('跳操'):
        label = f'跳操（{item}）'
    elif cat == '跳绳':
        label = '跳绳 ' + item.replace('次/分钟', ' 次/分')
    else:
        label = item
    e = {'cat': re.sub(r'(\S+?)(速度|\*|休息|通勤).*', r'\1', cat), 'label': label, 'coef': coef, 'src': f'表16 E{row}'}
    if isflat:
        e['flat'] = True
    if item == '每走一万步':
        e['perSteps'] = True
    cardio.append(e)
cardio.insert(next(i for i, c in enumerate(cardio) if c['cat'].startswith('户外骑行')),
              {'cat': '跑步', 'label': '跑步', 'coef': None, 'src': '表16 D94-E103', 'run': True})
# 心率法（方法一）用的是哪套系数
hr_flat = factor_of_row.get(14) == flat

# ---------- 表19 ----------
V19 = rows_of(wv['19日常食物营养率'], 1, 140)
def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
def pct(note, key):
    m = re.search(key + r'(\d+)%', note or '')
    return int(m.group(1)) / 100 if m else None
food = {'carb': [], 'protein': [], 'mixed': []}
section, cat = None, ''
for row in sorted(V19):
    d = V19[row]
    b = str(d.get('B', '')).strip()
    if b == '碳水': section = 'carb'; continue
    if b == '蛋白质': section = 'protein'; continue
    if b == '脂肪': section = None; continue
    if b in ('大类',) or section is None:
        continue
    if b:
        cat = b.replace('\n', '')
    name = str(d.get('C', '')).strip()
    if not name or name.startswith('【'):
        continue
    rate, note = d.get('D'), str(d.get('F', '') or '').strip()
    sec = 'mixed' if cat.startswith('含有蛋白质') else section
    item = {'cat': cat, 'n': name, 'src': f'表19 C{row}'}
    if num(rate) is not None:
        item['r'] = num(rate)
    else:
        item['fixed'] = str(rate)
    if sec == 'carb':
        item['gi'] = str(d.get('E', '') or '')
    else:
        item['pos'] = str(d.get('E', '') or '')
        c, p, f = pct(note, '碳水'), pct(note, '蛋白质'), pct(note, '脂肪')
        if c is not None: item['c'] = c
        if f is not None: item['f'] = f
    if note:
        item['note'] = note
    food[sec].append(item)

# ---------- 表26/27 拉伸图 ----------
z = zipfile.ZipFile(SRC)
wbx = z.read('xl/workbook.xml').decode()
rels_x = z.read('xl/_rels/workbook.xml.rels').decode()
rid2t = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels_x)}
rid2t.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels_x)})
# 每张图的肌肉名（看图读出的标题）和它对应的训练肌群；原表 26_03 和 26_04 是同一张肱三头肌，只留一张
TITLES = {
    ('26', 0): ('胸大肌 肩前束', ['mid_chest', 'low_chest', 'up_chest', 'front']), ('26', 1): ('背阔肌', ['pull', 'row', 'back']),
    ('26', 2): ('肱二头肌', ['bi']), ('26', 3): ('肱三头肌', ['tri']), ('26', 5): ('肩中束', ['side']), ('26', 6): ('肩后束', ['rear']),
    ('26', 7): ('上斜方肌', []), ('26', 8): ('中下斜方肌', ['row', 'rear']), ('26', 9): ('腹直肌', ['abs']), ('26', 10): ('竖脊肌', ['ham', 'row']),
    ('27', 0): ('股四头肌', ['quad', 'comp']), ('27', 1): ('腘绳肌', ['ham']), ('27', 2): ('臀大肌', ['glute', 'comp']), ('27', 3): ('腓肠肌 比目鱼肌', ['quad', 'comp']),
}
stretch = []
for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
    no = name[:2]
    if no not in ('26', '27'):
        continue
    sp = 'xl/' + rid2t[rid].replace('/xl/', '').lstrip('/')
    rp = posixpath.join(posixpath.dirname(sp), '_rels', posixpath.basename(sp) + '.rels')
    for dref in re.findall(r'Target="([^"]*drawing[^"]*)"', z.read(rp).decode()):
        dp = posixpath.normpath(posixpath.join(posixpath.dirname(sp), dref))
        drp = posixpath.join(posixpath.dirname(dp), '_rels', posixpath.basename(dp) + '.rels')
        rmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(drp).decode()))
        anchors = re.findall(r'<xdr:from><xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>.*?r:embed="([^"]+)"', z.read(dp).decode(), re.S)
        anchors.sort(key=lambda a: (int(a[1]), int(a[0])))
        for i, (col, row, emb) in enumerate(anchors):
            if (no, i) not in TITLES:
                continue
            ip = posixpath.normpath(posixpath.join(posixpath.dirname(dp), rmap[emb]))
            im = Image.open(io.BytesIO(z.read(ip))).convert('RGB')
            im.thumbnail((1000, 1000))
            fn = f'{no}_{i:02d}.jpg'
            im.save('www/img/stretch/' + fn, quality=72, optimize=True, progressive=True)
            t, groups = TITLES[(no, i)]
            stretch.append({'m': t, 'img': 'img/stretch/' + fn, 'groups': groups, 'src': f'表{no} 第{int(row) + 1}行{"左" if int(col) < 10 else "右"}图'})

js = '/* 自动生成：python3 tools/build-tables.py 套表.xlsx —— 表16、表19 全表与表26/27 拉伸图谱，逐行来自原表 */\n'
js += 'window.CARDIO = ' + json.dumps(cardio, ensure_ascii=False) + ';\n'
js += 'window.RUN_TABLE = ' + json.dumps(run_table) + ';\n'
js += 'window.CARDIO_WEIGHT_FACTOR = ' + json.dumps({'weights': WEIGHTS, 'normal': list(normal), 'flat': list(flat), 'hrFlat': hr_flat}) + ';\n'
js += 'window.FOOD19 = ' + json.dumps(food, ensure_ascii=False) + ';\n'
js += 'window.STRETCH = ' + json.dumps(stretch, ensure_ascii=False) + ';\n'
open('www/js/data-tables.js', 'w').write(js)
json.dump({'weights': WEIGHTS, 'rows': {str(k): {'coef': V16[k]['E'], 'item': str(V16[k].get('D', '')), 'values': v, 'flat': factor_of_row.get(k) == flat} for k, v in expect.items()}},
          open('test/fixtures/table16.json', 'w'), ensure_ascii=False)
print(f'表16 {len(cardio)} 项 + 跑步 {len(run_table)} 档；表19 碳水 {len(food["carb"])} 蛋白质 {len(food["protein"])} 混合 {len(food["mixed"])}；拉伸图 {len(stretch)} 张；心率法用{"平地" if hr_flat else "一般"}系数')
